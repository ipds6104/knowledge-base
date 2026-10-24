#!/usr/bin/env python3
"""
Kompilasi Dataset Murni Mentah (100% Raw) PODES 2026 Kabupaten Mempawah.
Mengonversi 31 tabel mentah JSON menjadi format Parquet (.parquet) dan Excel (.xlsx) murni
tanpa penggabungan/pengolahan sintetis (zero synthetic modifications).
"""

import json
import os
import shutil
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pyarrow as pa
import pyarrow.parquet as pq

BASE_DIR = Path(__file__).resolve().parents[1]
JSON_DIR = BASE_DIR / "scratch" / "raw_podes_json"
PARQUET_DIR = BASE_DIR / "data" / "podes-2026" / "parquet"
EXCEL_DIR = BASE_DIR / "data" / "podes-2026" / "excel"

EXCEL_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_EXCEL = EXCEL_DIR / "PODES_2026_Kabupaten_Mempawah.xlsx"

# Manifest tabel & pemetaan nama sheet Excel (Maksimal 31 karakter untuk Excel)
TABLE_SPECS = [
  # --- 1. ASET FASILITAS & SARANA PRASARANA RIIL (BY NAME & DESA) ---
  ("raw_aset_pendidikan_negeri", "Aset_Pendidikan_Negeri", "Daftar Aset Fisik Pendidikan Negeri per Desa"),
  ("raw_aset_pendidikan_swasta", "Aset_Pendidikan_Swasta", "Daftar Aset Fisik Pendidikan Swasta per Desa"),
  ("raw_aset_kesehatan", "Aset_Kesehatan", "Daftar Aset Fisik Fasilitas Kesehatan per Desa"),
  ("raw_aset_perbankan", "Aset_Perbankan", "Daftar Aset Fisik Bank & BPR per Desa"),
  ("raw_aset_ekonomi_pasar", "Aset_Ekonomi_Pasar", "Daftar Aset Fisik Pasar & Perdagangan per Desa"),

  # --- 2. TABULASI FASILITAS & POTENSI DESA (67 DESA) ---
  ("raw_tabulasi_desa_master", "Tab_Status_Desa_R301", "Tabulasi Status Administrasi Desa R301"),
  ("raw_tabulasi_kewilayahan_blok3", "Tab_Kewilayahan_Blok3", "Tabulasi Karakteristik Wilayah Blok 3"),
  ("raw_tabulasi_pendidikan_tk_ra", "Tab_Pendidikan_TK_RA", "Tabulasi Fasilitas TK dan RA"),
  ("raw_tabulasi_pendidikan_sd_mi", "Tab_Pendidikan_SD_MI", "Tabulasi Fasilitas SD dan MI"),
  ("raw_tabulasi_pendidikan_smp_mts", "Tab_Pendidikan_SMP_MTs", "Tabulasi Fasilitas SMP dan MTs"),
  ("raw_tabulasi_pendidikan_sma_smk", "Tab_Pendidikan_SMA_SMK", "Tabulasi Fasilitas SMA, SMK, dan MA"),
  ("raw_tabulasi_pendidikan_ptn", "Tab_Pendidikan_PTN", "Tabulasi Akademi & Perguruan Tinggi"),
  ("raw_tabulasi_pendidikan_daycare", "Tab_Daycare", "Tabulasi Tempat Penitipan Anak / Daycare"),
  ("raw_tabulasi_kesehatan_rs_balai", "Tab_Kesehatan_RS_Balai", "Tabulasi Rumah Sakit & Balai Pengobatan"),
  ("raw_tabulasi_kesehatan_puskesmas", "Tab_Kesehatan_Puskesmas", "Tabulasi Puskesmas & Pustu"),
  ("raw_tabulasi_kesehatan_klinik_praktek", "Tab_Kesehatan_Klinik", "Tabulasi Klinik & Praktek Dokter/Bidan"),
  ("raw_tabulasi_kesehatan_poskes_apotek", "Tab_Kesehatan_Poskesdes", "Tabulasi Poskesdes, Polindes, Apotek"),
  ("raw_tabulasi_ekonomi_pasar_swalayan", "Tab_Ekonomi_Pasar_Swalayan", "Tabulasi Pasar & Swalayan/Minimarket"),
  ("raw_tabulasi_ekonomi_restoran_toko", "Tab_Ekonomi_Restoran_Toko", "Tabulasi Restoran, Rumah Makan, Toko"),
  ("raw_tabulasi_ekonomi_bank", "Tab_Ekonomi_Bank_BPR", "Tabulasi Bank Umum & BPR"),
  ("raw_tabulasi_wilayah_laut_pesisir", "Tab_Laut_Pesisir", "Tabulasi Karakteristik Desa Laut/Pesisir"),
  ("raw_tabulasi_blok7_lainnya", "Tab_Blok7_Lainnya", "Tabulasi Variabel Tambahan Blok 7"),
  ("raw_tabulasi_penduduk_dtsen", "Tab_Penduduk_DTSEN", "Tabulasi Penduduk DTSEN per Kelompok Umur"),

  # --- 3. INDEKS KESULITAN GEOGRAFIS (IKG) (67 DESA) ---
  ("raw_ikg_desa_total", "IKG_Desa_Total", "Skor IKG 2025 vs 2026 dan Selisih"),
  ("raw_ikg_pelayanan_dasar_desa", "IKG_Pelayanan_Dasar", "IKG Dimensi Pelayanan Dasar"),
  ("raw_ikg_infrastruktur_desa", "IKG_Infrastruktur", "IKG Dimensi Infrastruktur"),
  ("raw_ikg_aksesibilitas_desa", "IKG_Aksesibilitas", "IKG Dimensi Aksesibilitas"),

  # --- 4. MONITORING PROGRES LAPANGAN ---
  ("raw_monitoring_desa", "Mon_Progres_Desa", "Monitoring Progres CAPI 67 Desa"),
  ("raw_monitoring_kecamatan", "Mon_Progres_Kecamatan", "Monitoring Progres CAPI 9 Kecamatan"),
  ("raw_monitoring_kabupaten", "Mon_Progres_Kabupaten", "Monitoring Progres CAPI Kabupaten"),
  ("raw_monitoring_pcl", "Mon_Progres_PCL", "Monitoring Alur Dokumen Petugas PCL")
]

# Styling BPS
HEADER_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ZEBRA_FILL = PatternFill(start_color="F7F9FC", end_color="F7F9FC", fill_type="solid")
THIN_SIDE = Side(border_style="thin", color="D9D9D9")
BORDER_CELL = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
BORDER_HEADER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=Side(border_style="medium", color="0D1E33"))

def style_worksheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_HEADER
    ws.row_dimensions[1].height = 28

    max_col = ws.max_column
    for row_idx in range(2, ws.max_row + 1):
        is_even = (row_idx % 2 == 0)
        row_fill = ZEBRA_FILL if is_even else PatternFill(fill_type=None)
        ws.row_dimensions[row_idx].height = 20

        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if row_fill.fill_type:
                cell.fill = row_fill
            cell.border = BORDER_CELL
            cell.font = Font(name="Calibri", size=10)

            val = cell.value
            header = str(ws.cell(row=1, column=col_idx).value or "")

            if isinstance(val, (int, float)):
                if "%" in header or "persen" in header.lower():
                    cell.number_format = "0.00%" if abs(val) <= 1.0 else "0.00"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif any(k in header.lower() for k in ["ikg", "dim.", "selisih"]):
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                if col_idx == 1 and header.lower() in ("no", "kode"):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif any(k in header.lower() for k in ["kode", "flag", "id", "status"]):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if cell.row == 1:
                val_str = val_str.replace("\n", " ")
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 60), 12)

def main():
    print("🚀 Mengompilasi 31 Tabel Raw PODES 2026...")

    # 1. Bersihkan direktori Parquet agar tidak ada tabel non-raw lama
    if PARQUET_DIR.exists():
        shutil.rmtree(PARQUET_DIR)
    PARQUET_DIR.mkdir(parents=True, exist_ok=True)

    # 2. Inisialisasi Workbook Excel
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Hapus default sheet

    parquet_count = 0
    excel_count = 0

    for key, sheet_title, desc in TABLE_SPECS:
        json_file = JSON_DIR / f"{key}.json"
        if not json_file.exists():
            print(f"⚠️ Berkas {key}.json tidak ditemukan, melewati...")
            continue

        with open(json_file, "r", encoding="utf-8") as f:
            jdata = json.load(f)

        records = jdata.get("data", [])
        colnames = jdata.get("colnames", [])

        if not colnames and records:
            colnames = list(records[0].keys())

        # A. Simpan ke Parquet (.parquet) murni
        # Normalisasi struktur data kolom
        pydict = {col: [r.get(col) for r in records] for col in colnames}
        arrow_table = pa.Table.from_pydict(pydict)
        pq_path = PARQUET_DIR / f"{key}.parquet"
        pq.write_table(arrow_table, pq_path, compression="snappy")
        parquet_count += 1

        # B. Simpan ke Excel Sheet tersendiri murni
        ws = wb.create_sheet(title=sheet_title)
        headers = ["No"] + colnames
        ws.append(headers)

        for idx, r in enumerate(records, 1):
            row_vals = [idx] + [r.get(col) for col in colnames]
            ws.append(row_vals)

        style_worksheet(ws)
        excel_count += 1
        print(f"  → [{excel_count}/31] {sheet_title:26} : {len(records):4} baris, {len(colnames):2} kolom (Parquet & Excel OK)")

    # Simpan workbook Excel
    wb.save(OUTPUT_EXCEL)
    excel_size_kb = os.path.getsize(OUTPUT_EXCEL) / 1024

    print(f"\n🎉 SELESAI SEMPURNA!")
    print(f"📁 Parquet: {parquet_count} berkas mentah murni di: {PARQUET_DIR}")
    print(f"📊 Excel  : 1 file terpadu ({excel_count} Sheet) di:")
    print(f"   {OUTPUT_EXCEL} ({excel_size_kb:.2f} KB)")

if __name__ == "__main__":
    main()
