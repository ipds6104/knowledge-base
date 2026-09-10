"""Legacy KCDA 2025 Data Extractor & Seeder Module.

Extracts historical and static tables from KCDA 2025 documents (DOCX & PDF)
and populates the upstream Google Sheets or local cache for KCDA 2026.
"""

import os
import json
import re
from pathlib import Path
from typing import Dict, Any, List, Optional
import docx
import pymupdf
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build

KEC_FILES = {
    'Anjongan': 'KCDA ANJONGAN_REV KONTEN TERAKHIRRR BGT.docx',
    'Jongkat': 'KCDA JONGKAT_REV KONTEN TERAKHIRRR BGT.docx',
    'Mempawah Hilir': 'KCDA MEMPAWAH HILIR_REV KONTEN TERAKHIRRR BGT.docx',
    'Sadaniang': 'KCDA SADANIANG_REV KONTEN TERAKHIRRR BGT.docx',
    'Segedong': 'KCDA SEGEDONG_REV KONTEN TERAKHIRRR BGT.docx',
    'Sungai Kunyit': 'KCDA SUNGAI KUNYIT_REV KONTEN TERAKHIRRR BGT.docx',
    'Sungai Pinyuh': 'KCDA SUNGAI PINYUH_REV KONTEN TERAKHIRRR BGT.docx',
    'Toho': 'KCDA TOHO_REV KONTEN TERAKHIRRR BGT.docx',
    'Mempawah Timur': 'KCDA MEMPAWAH TIMUR_REV KONTEN TERAKHIRRR BGT.pdf'
}

BASE_REF_DIR = Path('kegiatan/kecamatan-dalam-angka/referensi-2025')
CACHE_DIR = Path('data/kcda-2026/raw_tables')


def clean_str(val: Any) -> str:
    if val is None:
        return ""
    text = str(val).strip()
    text = re.sub(r'\s+', ' ', text)
    return text


def extract_docx_tables(fpath: Path) -> Dict[str, List[List[str]]]:
    """Ekstraksi tabel dari file docx KCDA 2025."""
    doc = docx.Document(fpath)
    extracted = {}

    for i, t in enumerate(doc.tables):
        c0 = t.rows[0].cells[0].text.strip() if t.rows and t.rows[0].cells else ''
        if c0 == 'Tabel':
            title = t.rows[0].cells[1].text.strip().replace('\n', ' ') if len(t.rows[0].cells) > 1 else ''
            dt = doc.tables[i + 1] if i + 1 < len(doc.tables) else None
            if not dt:
                continue

            t_lower = title.lower()
            # 1.2 Jarak ke Ibukota Kecamatan dan Kabupaten
            if 'jarak ke ibukota kecamatan' in t_lower and 'jarak_1_2' not in extracted:
                rows = []
                for r in dt.rows[2:]: # Skip 2 header rows
                    cells = [clean_str(c.text) for c in r.cells]
                    if cells and cells[0] and not any(k in cells[0].lower() for k in ['sumber', 'catatan']):
                        rows.append(cells)
                extracted['jarak_1_2'] = rows

            # 1.4 Jarak Kantor Camat ke Tempat Penting
            elif 'jarak kantor camat' in t_lower and 'jarak_1_4' not in extracted:
                rows = []
                for r in dt.rows[2:]:
                    cells = [clean_str(c.text) for c in r.cells]
                    if cells and len(cells) >= 3 and cells[1] and not any(k in cells[1].lower() for k in ['sumber', 'catatan']):
                        rows.append(cells)
                extracted['jarak_1_4'] = rows

            # 2.1.1 Rukun Tetangga dan Rukun Warga
            elif ('rukun warga' in t_lower or 'rukun tetangga' in t_lower) and 'rt_rw_2_1_1' not in extracted:
                rows = []
                for r in dt.rows[2:]:
                    cells = [clean_str(c.text) for c in r.cells]
                    if cells and cells[0] and not any(k in cells[0].lower() for k in ['sumber', 'catatan']):
                        rows.append(cells)
                extracted['rt_rw_2_1_1'] = rows

            # 2.1.2 Nama-Nama Camat
            elif ('camat yang pernah' in t_lower or 'nama-nama camat' in t_lower) and 'camat_2_1_2' not in extracted:
                rows = []
                for r in dt.rows[2:]:
                    cells = [clean_str(c.text) for c in r.cells]
                    if cells and len(cells) >= 3 and cells[1] and not any(k in cells[1].lower() for k in ['sumber', 'catatan']):
                        rows.append(cells)
                extracted['camat_2_1_2'] = rows

    return extracted


def extract_mempawah_timur_pdf(fpath: Path) -> Dict[str, List[List[str]]]:
    """Ekstraksi tabel dari PDF Mempawah Timur 2025 menggunakan PyMuPDF."""
    doc = pymupdf.open(fpath)
    extracted = {}

    # 1.2 Jarak ke Ibukota (Page 30 / index 29)
    p30 = doc[29]
    for tb in p30.find_tables().tables:
        rows = tb.extract()
        if len(rows) > 3:
            clean_rows = []
            for r in rows[2:]:
                c0 = clean_str(r[0])
                c1 = clean_str(r[1])
                c2 = clean_str(r[2]) if len(r) > 2 else ""
                if c0 and not any(k in c0.lower() for k in ['sumber', 'catatan']):
                    clean_rows.append([c0, c1, c2])
            extracted['jarak_1_2'] = clean_rows
            break

    # 1.4 Jarak Kantor Camat (Page 32 / index 31)
    for p_idx in range(30, 36):
        p = doc[p_idx]
        if 'Jarak Kantor Camat' in p.get_text():
            for tb in p.find_tables().tables:
                rows = tb.extract()
                if len(rows) > 2:
                    clean_rows = []
                    for r in rows[2:]:
                        c0 = clean_str(r[0])
                        c1 = clean_str(r[1])
                        c2 = clean_str(r[2]) if len(r) > 2 else ""
                        if c1 and not any(k in c1.lower() for k in ['sumber', 'catatan']):
                            clean_rows.append([c0, c1, c2])
                    extracted['jarak_1_4'] = clean_rows
                    break
            if 'jarak_1_4' in extracted:
                break

    # 2.1.1 RT / RW (Page 39 / index 38)
    for p_idx in range(35, 45):
        p = doc[p_idx]
        if 'Rukun Tetangga' in p.get_text():
            for tb in p.find_tables().tables:
                rows = tb.extract()
                if len(rows) > 3:
                    clean_rows = []
                    for r in rows[2:]:
                        c0 = clean_str(r[0])
                        c1 = clean_str(r[1])
                        c2 = clean_str(r[2]) if len(r) > 2 else ""
                        c3 = clean_str(r[3]) if len(r) > 3 else ""
                        if c0 and not any(k in c0.lower() for k in ['sumber', 'catatan']):
                            clean_rows.append([c0, c1, c2, c3])
                    extracted['rt_rw_2_1_1'] = clean_rows
                    break
            if 'rt_rw_2_1_1' in extracted:
                break

    # 2.1.2 Camat (Page 41 / index 40)
    for p_idx in range(38, 48):
        p = doc[p_idx]
        if 'Nama-Nama Camat' in p.get_text():
            for tb in p.find_tables().tables:
                rows = tb.extract()
                if len(rows) > 3:
                    clean_rows = []
                    for r in rows[2:]:
                        c0 = clean_str(r[0])
                        c1 = clean_str(r[1])
                        c2 = clean_str(r[2]) if len(r) > 2 else ""
                        if c1 and not any(k in c1.lower() for k in ['sumber', 'catatan']):
                            clean_rows.append([c0, c1, c2])
                    extracted['camat_2_1_2'] = clean_rows
                    break
            if 'camat_2_1_2' in extracted:
                break

    return extracted


def seed_legacy_data_to_cache(push_gsheet: bool = False) -> Dict[str, Any]:
    """
    Mengekstrak data statis dari 9 publikasi KCDA 2025 dan menyuntikkannya ke cache JSON lokal
    serta secara opsional mengunggah ke Google Sheets hulu.
    """
    print("🌾 Memulai ekstraksi data historis/statis dari naskah KCDA 2025 (9 Kecamatan)...")

    all_extracted = {}
    for kec, fname in KEC_FILES.items():
        fpath = BASE_REF_DIR / fname
        if not fpath.exists():
            print(f"   ⚠️ Berkas {fpath} tidak ditemukan, lewati {kec}.")
            continue

        if fname.endswith('.docx'):
            data = extract_docx_tables(fpath)
        else:
            data = extract_mempawah_timur_pdf(fpath)

        all_extracted[kec] = data
        stats = ", ".join([f"{k}: {len(v)} baris" for k, v in data.items()])
        print(f"   ✅ {kec:<15} ({stats})")

    # Injeksi ke Cache JSON
    target_mappings = [
        ('jarak_1_2', 'tabel_1_2_jarak_ke_ibukota_kecamatan_dan.json'),
        ('jarak_1_4', 'tabel_1_4_jarak_kantor_camat_xxx_dengan.json'),
        ('rt_rw_2_1_1', 'tabel_2_1_1_jumlah_rukun_warga__rw__dan_ru.json'),
        ('camat_2_1_2', 'tabel_2_1_2_nama_nama_camat_yang_pernah_ma.json')
    ]

    print("\n💾 Menyuntikkan data ekstraksi ke file cache JSON KCDA 2026...")
    for key, json_name in target_mappings:
        json_path = CACHE_DIR / json_name
        if not json_path.exists():
            continue

        with open(json_path, encoding='utf-8') as f:
            t_json = json.load(f)

        tabs = t_json.get('tabs', {})
        updated_tabs = 0
        for kec, ext_data in all_extracted.items():
            if key not in ext_data:
                continue

            extracted_rows = ext_data[key]
            # Match tab by kec name
            matched_tab = None
            for tab_name in tabs.keys():
                if kec.lower() in tab_name.lower():
                    matched_tab = tab_name
                    break

            if matched_tab:
                orig_rows = tabs[matched_tab].get('rows', [])
                header_rows = orig_rows[:3] if len(orig_rows) >= 3 else []
                # Gabungkan header dengan data baru
                tabs[matched_tab]['rows'] = header_rows + extracted_rows
                updated_tabs += 1

        t_json['tabs'] = tabs
        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(t_json, f, indent=2, ensure_ascii=False)

        print(f"   ✨ Berhasil memperbarui {json_name} ({updated_tabs} kecamatan)")

    if push_gsheet:
        print("\n☁️  Mendorong data ke Google Sheets hulu via Sheets API...")
        push_seeded_data_to_gsheet(target_mappings, all_extracted)

    return all_extracted


def push_seeded_data_to_gsheet(target_mappings, all_extracted):
    """Menuliskan data ke Google Sheets menggunakan kredensial token.json."""
    if not os.path.exists('token.json'):
        print("   ❌ token.json tidak ditemukan, lewati pengunggahan Google Sheets.")
        return

    creds = Credentials.from_authorized_user_file('token.json')
    service = build('sheets', 'v4', credentials=creds)

    for key, json_name in target_mappings:
        json_path = CACHE_DIR / json_name
        with open(json_path, encoding='utf-8') as f:
            t_json = json.load(f)

        sid = t_json.get('sheet_id')
        if not sid:
            continue

        meta = service.spreadsheets().get(spreadsheetId=sid).execute()
        sheet_titles = {s['properties']['title'].lower(): s['properties']['title'] for s in meta['sheets']}

        for kec, ext_data in all_extracted.items():
            if key not in ext_data:
                continue

            matched_title = None
            for s_lower, orig_title in sheet_titles.items():
                if kec.lower() in s_lower:
                    matched_title = orig_title
                    break

            if not matched_title:
                continue

            rows_to_write = ext_data[key]
            if not rows_to_write:
                continue

            range_to_update = f"'{matched_title}'!A4:E{3 + len(rows_to_write)}"
            body = {'values': rows_to_write}

            try:
                service.spreadsheets().values().update(
                    spreadsheetId=sid,
                    range=range_to_update,
                    valueInputOption='USER_ENTERED',
                    body=body
                ).execute()
                print(f"      📤 {t_json['no']} [{kec}] -> Google Sheet terupdate ({len(rows_to_write)} baris)")
            except Exception as e:
                print(f"      ⚠️ Gagal update {kec}: {e}")
